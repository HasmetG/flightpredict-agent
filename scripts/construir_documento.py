"""Genera el documento de la entrega (DOCX editable) leyendo las cifras de results/*.csv, para que el texto coincida con los logs.
Uso:  python scripts/construir_documento.py   (requiere python-docx)  ->  docs/FlightPredict_Agent_Entrega_Final.docx
Luego: abrir en Word / LibreOffice, completar los campos resaltados en amarillo y exportar a PDF."""
import gzip
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
R, D = ROOT / "results", ROOT / "docs"
OUT = D / "FlightPredict_Agent_Entrega_Final.docx"

tabla = pd.read_csv(R / "tabla_resultados.csv")
esc = pd.read_csv(R / "resumen_por_escenario.csv")
det = pd.read_csv(R / "detalle_por_corrida.csv")
res = json.load(open(R / "resumen.json"))
sens = pd.read_csv(R / "sensibilidad_utilidad_validacion.csv")
thr = res["umbrales"]


def pc(x, d=1):  # 0.7403 -> "74,0 %"
    return f"{x * 100:.{d}f}".replace(".", ",") + " %"


def num(x, d=1):
    return f"{x:.{d}f}".replace(".", ",")


def miles(n):
    return f"{int(n):,}".replace(",", ".")


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(2.0))
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(5)
for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
    h = doc.styles[name]
    h.font.name, h.font.size, h.font.bold = "Calibri", Pt(size), True
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(10), Pt(4)


def para(text="", bold=False, italic=False, size=None, align=None, after=None):
    p = doc.add_paragraph()
    add_runs(p, text, bold, italic, size)
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def add_runs(p, text, bold=False, italic=False, size=None):
    """Texto con **negrita** y [[campo a completar]] (resaltado amarillo)."""
    import re
    for tok in re.split(r"(\*\*.+?\*\*|\[\[.+?\]\])", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("[["):
            r = p.add_run("[" + tok[2:-2] + "]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = p.add_run(tok); r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    add_runs(p, text)
    p.paragraph_format.space_after = Pt(2)
    return p


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(header, rows, widths, size=8, hdr="D9E2F3"):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        add_runs(c.paragraphs[0], h, bold=True, size=size)
        shade(c, hdr)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], str(v), size=size)
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(widths) / 2.54 * 1440)))
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
    # repetir encabezado en cada pagina
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, width_cm, caption):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    c = para(caption, italic=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    return c


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---- pie de pagina con numero de pagina --------------------------------------------------------------
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("UTB — Inteligencia Artificial — FlightPredict-Agent · Página "); r.font.size = Pt(8)
r2 = fp.add_run(); r2.font.size = Pt(8)
for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if kind:
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
    else:
        e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
    r2._r.append(e)

# =====================================================================================================
# PORTADA
# =====================================================================================================
for _ in range(4):
    para("")
para("UNIVERSIDAD TECNOLÓGICA DE BOLÍVAR", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Ingeniería de Sistemas y Computación — Inteligencia Artificial", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
para("")
para("Diseño e implementación de un agente inteligente", bold=True, size=20, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Documento de diseño ajustado, desarrollo del agente y pruebas contra las medidas de desempeño", size=11.5,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("")
para("FlightPredict-Agent — Agente de predicción de retrasos de vuelo", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(2):
    para("")
para("Presentado por:", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Hasmet Gonzalez Perez — T00070395", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
para("Dylan Andrés Ecker López — T00078024", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
para("")
para("Docente: Edwin Alexander Puertas Del Castillo", align=WD_ALIGN_PARAGRAPH.CENTER)
para("Cartagena de Indias, 17 de octubre de 2026", align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(3):
    para("")
para("Contenido", bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
for line in ("0. Tabla de cambios respecto a la entrega anterior · Fase 1. Selección justificada del agente (matriz de decisión)",
             "1–8. Diseño del agente (corregido) · 9. Desarrollo del agente (arquitectura, tecnologías, entradas, salidas, "
             "diagramas de secuencia, diferencias, pruebas, limitaciones, enlaces) · 10. Conclusiones · Referencias · Anexo A"):
    para(line, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
page_break()

# =====================================================================================================
# 0. TABLA DE CAMBIOS
# =====================================================================================================
doc.add_heading("0. Tabla de cambios respecto a la entrega anterior (Tarea 2)", 1)
para("Las modificaciones se hicieron al implementar el agente y al contrastar el diseño con las exigencias de la Tarea 3. "
     "La fila final recoge las observaciones escritas del docente a la Tarea 2.")
ch = [
    ("Portada", "Un documento por integrante, con un solo autor cada uno.",
     "Un único documento de la pareja, que corresponde al agente elegido.",
     "La entrega es una sola por pareja; la elección se justifica en la Fase 1."),
    ("Nueva sección (Fase 1)", "No existía.", "Matriz de decisión ponderada, análisis de sensibilidad y justificación (antes del diseño).",
     "Exigido por la actividad: la elección debe poder reconstruirse."),
    ("2. Tipo de agente", "Comparaba la probabilidad con «un umbral definido» sin explicitar la función de utilidad.",
     "Se define U(alertar) = P·β·d − (1−P)·C_FA y se explica que el umbral es el punto de equilibrio P* = C_FA/(C_FA+β·d), que varía con la duración esperada.",
     "Un agente basado en utilidad debe comparar alternativas con una función de utilidad explícita y verificable en el código."),
    ("3. Tipos de ambiente (dinamismo)", "El agente «debe recalcular» la probabilidad a medida que se acerca la hora del vuelo.",
     "El agente emite una predicción por vuelo a T − 180 min; el recálculo periódico queda como trabajo futuro.",
     "Se implementó solo la predicción a las 3 h; el recálculo no se pudo probar en el tiempo disponible."),
    ("5. Entradas y sensores", "Historial BTS con 32 variables; clima desde una API externa.",
     "Se agregan identificador de aeronave, número de tramo, llegada programada del tramo previo y reloj del sistema. Las fuentes se simulan (ver 9.3); se usan 15 características.",
     "Sin el identificador de aeronave y la hora actual no se pueden calcular el retraso en cadena ni la antelación."),
    ("6. Salidas y actuadores", "Tres actuadores (alerta, sugerencia, registro del resultado).",
     "Cuarto actuador: «marcar para revisión manual» ante entradas inválidas o ventana menor a 3 h.",
     "La sección 8 ya hablaba de marcar entradas defectuosas, pero no existía el actuador."),
    ("7. Medidas de desempeño", "Una sola medida compuesta (exactitud ≥ 70 % con 3 h de antelación), con evaluación continua y revisión trimestral.",
     "Cuatro medidas con fórmula y umbral: M1 exactitud, M2 antelación, M3 recall de vuelos retrasados y M4 manejo de fallas; 13 escenarios × 5 repeticiones.",
     "Con ≈ 20 % de vuelos retrasados, predecir «nunca hay retraso» ya da ≈ 80 % de exactitud; hacía falta una medida de sensibilidad. La revisión trimestral no es verificable en la entrega."),
    ("8. Memoria y razonamiento", "«Consultar el historial de vuelos similares» sin concretar cómo.",
     "Se concreta como tasa histórica de retraso por aeropuerto de origen y franja horaria, incluida en la alerta como contexto; la memoria de corto plazo guarda el último retraso conocido por aeronave.",
     "Se necesita una definición implementable y comprobable de cada componente."),
    ("9 → 10. Conclusiones", "Conclusiones de viabilidad conceptual.",
     "Pasa a ser la sección 10 y se actualiza con los resultados reales (incluida la medida M3 incumplida).",
     "Se inserta la sección 9 «Desarrollo del agente» antes de las conclusiones."),
    ("Nueva sección 9", "No existía.", "Arquitectura, tecnologías, entradas, salidas, diagramas de secuencia UML, diferencias, pruebas, limitaciones y enlaces.",
     "Exigido por la Tarea 3."),
    ("Referencias", "Russell y Norvig, IBM, BTS, PLOS ONE.", "Se agregan scikit-learn (Pedregosa et al., 2011) y CUAD (Hendrycks et al., 2021, usado en la matriz).",
     "Tecnología utilizada y fuente de la comparación."),
    ("Observaciones del docente", "—",
     "[[COMPLETAR: una fila por cada observación recibida en la Tarea 2 y la corrección aplicada]]", "[[COMPLETAR]]"),
]
table(["Sección", "Qué decía", "Qué dice ahora", "Por qué cambió"], ch, [2.6, 4.2, 5.8, 4.9])
page_break()

# =====================================================================================================
# FASE 1
# =====================================================================================================
doc.add_heading("Fase 1. Selección justificada del agente", 1)
para("Se comparan los dos agentes diseñados por la pareja: **Agente A — ClauseCheck-Agent** (revisión automatizada de contratos, basado en modelo con "
     "aprendizaje) y **Agente B — FlightPredict-Agent** (predicción de retrasos de vuelo, basado en utilidad con aprendizaje y modelo). "
     "Los criterios y los pesos se fijaron antes de puntuar.")
doc.add_heading("Criterios y pesos", 3)
table(["Criterio", "Peso", "Cómo se mide", "Por qué ese peso"], [
    ("Viabilidad técnica", "25 %", "Componentes IBM que el equipo sabe implementar / 7", "Es lo que más condiciona si el agente llega a existir en el tiempo de la entrega."),
    ("Disponibilidad de datos o entorno", "20 %", "Fuente identificada, simulador o dataset concreto", "Sin entorno y datos no hay percepciones que probar; es importante, pero ambos agentes pueden simular el suyo."),
    ("Medibilidad del desempeño", "20 %", "Medidas con fórmula, umbral y verdad terreno objetiva", "La actividad se califica con pruebas contra las medidas; por eso pesa igual que los datos."),
    ("Complejidad de IA", "15 %", "Tipo de agente según Russell y Norvig (reflejo simple → utilidad / aprendizaje)", "Debe haber razonamiento real, pero la actividad no es un concurso de sofisticación."),
    ("Costo y recursos", "10 %", "Dinero, hardware, APIs pagas o cuotas", "Importa al ser trabajo estudiantil sin presupuesto, pero no debe decidir por sí solo."),
    ("Valor del problema", "10 %", "Usuario y beneficio concretos", "Aporta sentido al proyecto, pero ambos problemas tienen usuario claro."),
    ("**Total**", "**100 %**", "", ""),
], [3.6, 1.4, 5.5, 7.0])
doc.add_heading("Matriz de decisión ponderada", 3)
table(["Criterio", "Peso (w)", "A: pts", "A: evidencia (ClauseCheck)", "B: pts", "B: evidencia (FlightPredict)"], [
    ("Viabilidad técnica", "25 %", "3", "El pipeline (extraer texto, segmentar cláusulas, clasificar con un modelo de lenguaje, registrar correcciones del abogado) es construible, pero exige PDF/OCR, un modelo de lenguaje y anotaciones; se estiman 5 de 7 componentes implementables sin infraestructura jurídica real (71 %).",
     "4", "Todo el flujo se resuelve con Python y scikit-learn en CPU; se estiman 6 de 7 componentes implementables (86 %); solo la integración con el sistema real de la aerolínea se simula."),
    ("Disponibilidad de datos o entorno", "20 %", "4", "CUAD (conjunto público con contratos anotados por abogados) permite validar; el gestor documental se simula con archivos y la retroalimentación del abogado con las etiquetas del conjunto.",
     "4", "El entorno se simula por completo (clima, rotaciones, retrasos en cadena) y el historial BTS es público; unirlo con clima histórico exigiría otra fuente."),
    ("Medibilidad del desempeño", "20 %", "3", "Una medida con fórmula y umbral (exactitud ≥ 80 % en 50 contratos), pero la verdad terreno depende de un abogado y «exactitud» es ambigua (por cláusula, categoría o fragmento).",
     "4", "La verdad terreno es objetiva y automática (el vuelo salió o no ≥ 15 min tarde); medida con fórmula y umbral. Debilidad: la exactitud sola es engañosa con ≈ 20 % de retrasos."),
    ("Complejidad de IA", "15 %", "4", "Agente basado en modelo con componente de aprendizaje: comprensión semántica de texto y ajuste con correcciones.",
     "5", "Basado en utilidad con aprendizaje y modelo: estima probabilidad y duración, compara utilidades y mantiene estado de la aeronave (nivel superior en la escala de Russell y Norvig)."),
    ("Costo y recursos", "10 %", "2", "Requiere API de un modelo de lenguaje de pago (varias pasadas sobre 50 contratos largos) o GPU para ajuste fino, y OCR para escaneados.",
     "5", "Costo de 0 COP: datos simulados o públicos, scikit-learn en CPU, sin APIs de pago."),
    ("Valor del problema", "10 %", "4", "Áreas jurídicas: reduce la lectura previa de cada contrato; no sustituye el criterio del abogado.",
     "4", "Equipos de operaciones de aerolíneas: anticipar retrasos en cadena y reasignar tripulaciones y puertas."),
    ("**Puntaje ponderado S**", "**100 %**", "**3,35**", "S(A) = 0,25·3 + 0,20·4 + 0,20·3 + 0,15·4 + 0,10·2 + 0,10·4",
     "**4,25**", "S(B) = 0,25·4 + 0,20·4 + 0,20·4 + 0,15·5 + 0,10·5 + 0,10·4"),
], [2.6, 1.2, 0.9, 5.2, 0.9, 5.7], size=7.5)
para("Los puntajes de viabilidad son una estimación del equipo con base en lo descrito en cada diseño; quien aplique los mismos pesos y puntajes obtiene S(A) = 3,35 y S(B) = 4,25.", italic=True, size=9)
doc.add_heading("Análisis de sensibilidad", 3)
para("El criterio más influyente en la diferencia entre agentes es Costo y recursos (w × |pB − pA| = 0,10 × 3 = 0,30; le sigue Viabilidad técnica con 0,25). "
     "Se cambió su peso en ±10 puntos y los demás pesos se reescalaron proporcionalmente para que sumen 100 %; se repitió con Viabilidad técnica, el criterio de mayor peso.")
table(["Escenario", "Pesos (Viab / Datos / Medib / Compl / Costo / Valor)", "S(A) ClauseCheck", "S(B) FlightPredict", "Decisión"], [
    ("Base", "25,0 / 20,0 / 20,0 / 15,0 / 10,0 / 10,0", "3,35", "4,25", "B"),
    ("Costo y recursos −10 pts (0 %)", "27,8 / 22,2 / 22,2 / 16,7 / 0,0 / 11,1", "3,50", "4,17", "B"),
    ("Costo y recursos +10 pts (20 %)", "22,2 / 17,8 / 17,8 / 13,3 / 20,0 / 8,9", "3,20", "4,33", "B"),
    ("Viabilidad técnica −10 pts (15 %)", "15,0 / 22,7 / 22,7 / 17,0 / 11,3 / 11,3", "3,40", "4,28", "B"),
    ("Viabilidad técnica +10 pts (35 %)", "35,0 / 17,3 / 17,3 / 13,0 / 8,7 / 8,7", "3,30", "4,22", "B"),
], [4.4, 6.2, 2.4, 2.6, 1.6], size=8)
para("La decisión no cambia en ninguno de los cuatro escenarios. Prueba extrema adicional: aun si ClauseCheck-Agent obtuviera 5 en Costo y recursos, S(A) sería 3,65 frente a 4,25 de B.")
doc.add_heading("Justificación de la elección", 3)
para("El equipo elige **FlightPredict-Agent** con un puntaje ponderado de 4,25 sobre 5 frente a 3,35 de ClauseCheck-Agent, una diferencia de 0,90 puntos. "
     "La ventaja viene de cuatro criterios: costo (el agente de vuelos corre en CPU sin APIs de pago), viabilidad técnica (el flujo completo se resuelve con Python y scikit-learn), "
     "medibilidad (la verdad terreno es un hecho objetivo, no la opinión de un abogado) y complejidad (el agente de utilidad con aprendizaje está más arriba en la escala de Russell y Norvig). "
     "Empatan en disponibilidad de datos y en valor del problema. La sensibilidad muestra que la elección es estable: al mover ±10 puntos los pesos de Costo o de Viabilidad, B sigue ganando por al menos 0,67 puntos.")
para("**Qué se pierde al descartar ClauseCheck-Agent:** la oportunidad de trabajar procesamiento de lenguaje natural sobre texto jurídico, un dominio con contratos anotados por expertos (CUAD) "
     "y un ciclo de retroalimentación humana más natural, y un problema con mayor carga ética (apoyo a la decisión de un abogado). A cambio, el equipo gana un agente que puede probarse de forma completa, "
     "repetible y con verdad terreno objetiva, que es lo que exige esta actividad.")
page_break()

# =====================================================================================================
# 1-8 DISEÑO
# =====================================================================================================
doc.add_heading("1. Introducción y planteamiento del problema", 1)
para("Los retrasos en la operación aérea representan uno de los problemas más costosos y disruptivos para las aerolíneas, dado que afectan la asignación de tripulaciones, la disponibilidad de puertas de embarque y la experiencia de los pasajeros. "
     "Cuando un vuelo se retrasa, el efecto suele propagarse en cadena hacia los siguientes vuelos programados con la misma aeronave, generando un impacto operativo que crece con el tiempo si no se anticipa oportunamente. "
     "Actualmente, muchas aerolíneas reaccionan a los retrasos una vez que estos ya se han materializado, lo que limita su capacidad de reorganizar recursos de manera proactiva.")
para("En este contexto, se propone FlightPredict-Agent, un agente inteligente orientado a estimar, con al menos tres horas de anticipación, la probabilidad de retraso de cada vuelo programado de una aerolínea piloto. "
     "Un agente inteligente constituye una solución pertinente porque permite integrar de manera autónoma múltiples fuentes de información heterogénea (meteorología, historial operativo, estado de la aeronave) "
     "y traducirlas en una recomendación accionable para el equipo de operaciones, sin depender exclusivamente del juicio manual de un analista humano bajo presión de tiempo.")

doc.add_heading("2. Tipo de agente inteligente", 1)
para("Según la clasificación de Russell y Norvig (Artificial Intelligence: A Modern Approach), FlightPredict-Agent se clasifica principalmente como un **agente basado en utilidad con capacidad de aprendizaje**. "
     "Es basado en utilidad porque no se limita a decidir si existe riesgo: calcula una probabilidad P de retraso y una duración esperada d, y compara la utilidad esperada de las dos acciones posibles. "
     "Función de utilidad: **U(alertar) = P·β·d − (1 − P)·C_FA**, con U(no alertar) = 0, donde β = 0,05 es el valor de anticipar cada minuto de retraso y C_FA = 1 el costo de una falsa alarma. "
     "El agente alerta cuando U(alertar) > 0, equivalente a superar el umbral P* = C_FA/(C_FA + β·d); como P* depende de d, un retraso esperado más largo justifica alertar con menos certeza "
     "(por ejemplo, P* = 0,50 si d = 20 min y P* = 0,25 si d = 60 min).")
para("Incorpora un componente de aprendizaje, dado que el modelo se entrena con el historial de vuelos y se retroalimenta con el resultado real de cada vuelo, y un componente basado en modelo, "
     "pues conserva una representación interna del estado de cada aeronave (el último retraso conocido) para inferir cómo ese estado, no observable directamente, afecta el vuelo actual.")

doc.add_heading("3. Tipos de ambiente", 1)
table(["Propiedad", "Clasificación", "Justificación"], [
    ("Observabilidad", "Parcialmente observable", "El agente solo percibe un subconjunto de variables (ruta, aerolínea, hora, pronóstico y último estado conocido de la aeronave); no ve la disponibilidad real de personal de tierra ni incidentes no reportados. En la implementación, el retraso del tramo previo solo se conoce si ese vuelo ya salió a T − 180 min."),
    ("Determinismo", "Estocástico", "La ocurrencia y la duración de un retraso dependen de factores probabilísticos (clima, congestión, tripulación); por eso el agente estima una probabilidad y no un resultado garantizado."),
    ("Episodicidad", "Secuencial", "La predicción de un vuelo depende del estado del vuelo anterior de la misma aeronave dentro de la jornada, de modo que las predicciones no son independientes."),
    ("Dinamismo", "Dinámico", "Clima, aeropuertos y aeronaves cambian mientras el agente delibera. Ajuste: la implementación emite una predicción por vuelo a T − 180 min; el recálculo periódico queda como trabajo futuro."),
    ("Discretización", "Mixto, con predominio continuo", "Hora, duración del retraso y magnitudes meteorológicas son continuas; aeropuerto, aerolínea y tramo son discretos."),
    ("Número de agentes", "Agente único", "Una sola entidad autónoma genera la predicción y la alerta; el sistema operativo y el equipo de operaciones forman parte del entorno."),
], [3.0, 3.6, 10.9], size=8.5)

doc.add_heading("4. Descripción del entorno", 1)
para("El agente opera dentro del contexto de una aerolínea piloto que sirve como caso de validación inicial. Los actores son el equipo de operaciones (reasigna tripulaciones y puertas), el sistema operativo de la aerolínea "
     "(con el cual el agente se integra para emitir alertas) y, de forma indirecta, los pasajeros. Entre las restricciones del dominio están la calidad de los pronósticos meteorológicos, la dependencia de fuentes históricas como el "
     "Bureau of Transportation Statistics (BTS) y la necesidad de predecir con al menos tres horas de antelación. En condiciones normales, el agente procesa vuelos con información estable; en condiciones excepcionales "
     "(tormentas severas, huelgas, fallas en cadena, registros incompletos o tardíos), el agente debe generalizar o, si no puede predecir, marcar el vuelo para revisión manual. "
     "Para la entrega, este entorno se reproduce con un simulador estocástico (sección 9).")

doc.add_heading("5. Entradas y sensores", 1)
for b in ("**Aeropuerto de origen y destino:** datos categóricos del sistema de programación de vuelos.",
          "**Aerolínea:** dato categórico del sistema de programación.",
          "**Hora planeada del vuelo:** marca de tiempo del itinerario operativo.",
          "**Identificador de aeronave, número de tramo y llegada programada del tramo previo (agregado):** permiten calcular la holgura y el retraso en cadena.",
          "**Reloj del sistema (agregado):** permite calcular la antelación de cada predicción.",
          "**Condiciones meteorológicas previstas:** visibilidad, viento, precipitación y temperatura, captadas de un servicio meteorológico que actúa como sensor.",
          "**Estado del vuelo anterior de la misma aeronave:** en la implementación, el último retraso conocido en minutos, entregado por el sistema operativo cuando el vuelo sale.",
          "**Historial de vuelos y retrasos:** fuente de entrenamiento. Diseño: 32 variables del BTS. Implementación: historial simulado con 15 características."):
    bullet(b)

doc.add_heading("6. Salidas y actuadores", 1)
for b in ("**Alerta al equipo de operaciones:** notificación estructurada al sistema operativo cuando U(alertar) > 0.",
          "**Sugerencia de reasignación proactiva:** ajuste de puerta, o de tripulación y puerta, hacia el sistema de gestión de recursos.",
          "**Registro del resultado real del vuelo:** escritura en el historial del agente (predicción + resultado) para el reentrenamiento.",
          "**Marcar para revisión manual (agregado):** señal al sistema operativo cuando el registro es inválido o la ventana es menor a tres horas; no se emite predicción."):
    bullet(b)

doc.add_heading("7. Medidas de desempeño (SMART)", 1)
para("La medida original (exactitud ≥ 70 % con tres horas de antelación) se descompone en cuatro medidas medibles, cada una con fórmula y umbral fijados antes de ejecutar las pruebas "
     f"(config/thresholds.json, fijado el {thr['fijado_utc'].replace('T', ' ').replace('Z', ' UTC')}). Un vuelo se considera retrasado si sale 15 minutos o más después de lo programado.")
table(["Medida", "Específica", "Medible (fórmula)", "Umbral", "Alcanzable / Relevante", "Temporal"], [
    ("M1 Exactitud", "Clasificar vuelos retrasados vs. no retrasados", "(VP + VN) / total", "≥ 70 % (media global)", "La literatura reporta niveles comparables; sin exactitud mínima las alertas pierden credibilidad.", "50 corridas (10 escenarios × 5 rep.)"),
    ("M2 Antelación", "Emitir cada predicción con ≥ 3 h de antelación", "predicciones con antelación ≥ 180 min / predicciones emitidas", "100 %", "Se controla con el disparador T − 180; sin ventana no hay reasignación posible.", "65 corridas (13 escenarios × 5 rep.)"),
    ("M3 Recall", "Detectar los vuelos que realmente se retrasan", "VP / (VP + FN)", "≥ 60 % (media global)", "Medida nueva: con ≈ 20 % de retrasos, la exactitud sola se logra sin detectar casi ninguno.", "50 corridas (S01–S10)"),
    ("M4 Manejo de fallas", "No fallar ante entradas inválidas, ventana < 3 h o sensor sin dato", "fallas manejadas / fallas inyectadas, sin excepciones", "100 % por escenario", "Un agente operativo no puede detenerse por un dato defectuoso.", "15 corridas (S11–S13)"),
], [2.0, 3.0, 3.3, 2.2, 4.4, 2.6], size=8)

doc.add_heading("8. Componentes de la arquitectura de IA agéntica (marco IBM)", 1)
para("Fuente: IBM, «Components of AI agents» — https://www.ibm.com/mx-es/think/topics/components-of-ai-agents. Cómo implementa el agente cada uno de los siete componentes (el módulo de código correspondiente se indica en 9.1):", size=9.5)
for t_, x in (
    ("Percepción y manejo de entradas.", "Para cada vuelo, el agente recibe ruta, aerolínea, hora planeada, pronóstico y estado de la aeronave. Los datos se validan antes de ingresar al modelo: un pronóstico ausente o fuera de rango físico se imputa con la climatología y se marca; un registro inválido se marca para revisión manual, evitando que una entrada defectuosa distorsione la estimación."),
    ("Planificación y descomposición de tareas.", "A las tres horas de la salida, el agente descompone la tarea en cinco subtareas secuenciales: recopilar y validar entradas; consultar el historial de vuelos similares (tasa histórica de retraso por aeropuerto de origen y franja horaria); estimar probabilidad y duración con el modelo; comparar utilidades; y emitir alerta y sugerencia. El plan queda registrado en el log de cada decisión, lo que permite auditar la predicción."),
    ("Memoria.", "Memoria de largo plazo: historial de vuelos y resultados reales, climatología por aeropuerto y mes (para imputar) y tasas históricas; sirve de base para entrenar y reentrenar el modelo. Memoria de corto plazo: último retraso conocido de cada aeronave durante la jornada, necesaria para anticipar retrasos en cadena."),
    ("Razonamiento y toma de decisiones.", "Un modelo de aprendizaje automático estima la probabilidad de retraso y su duración esperada incorporando el retraso en cadena; la decisión de alertar se toma comparando las utilidades esperadas (sección 2), no con reglas fijas."),
    ("Llamada de herramientas y acción.", "Cuando U(alertar) > 0, el agente invoca la conexión con el sistema operativo de la aerolínea para enviar la alerta y la sugerencia; para obtener el pronóstico consulta un servicio meteorológico. En la implementación ambos servicios están simulados."),
    ("Comunicación.", "La alerta es un mensaje estructurado con identificador del vuelo, probabilidad, duración esperada, sugerencia de reasignación y tasa histórica de vuelos similares, dentro de la ventana de tres horas."),
    ("Aprendizaje y adaptación.", "Al salir cada vuelo, el resultado real se compara con la predicción emitida y se acumula como retroalimentación; el agente reentrena el modelo con el historial y las muestras nuevas (que pesan más), para adaptarse a cambios de patrón."),
):
    para(f"**{t_}** {x}")
page_break()

# =====================================================================================================
# 9 DESARROLLO
# =====================================================================================================
doc.add_heading("9. Desarrollo del agente", 1)
para("El agente implementado corresponde al diseño de las secciones 1 a 8: mismo tipo (basado en utilidad con modelo y aprendizaje), mismos sensores y actuadores, y un entorno simulado que respeta las propiedades de la sección 3 "
     "(parcialmente observable, estocástico, secuencial, dinámico, mixto, agente único). Las diferencias se documentan en 9.6. El ciclo percepción → decisión → acción está en src/flightpredict/agent.py (métodos perceive, decide, act y step).")
doc.add_heading("9.1 Arquitectura", 2)
figure(D / "diagramas" / "arquitectura.png", 17.0,
       "Figura 1. Diagrama de componentes: sensores y servicios externos (simulados), módulos del agente y flujo de datos. Línea punteada: registro (log); línea discontinua: ruta de excepción.")
table(["Componente IBM", "Módulo de código (src/flightpredict)", "Estado", "Comentario"], [
    ("1. Percepción y manejo de entradas", "perception.py: perceive(), extract_features()", "Implementado", "Valida campos, rangos y antelación; imputa clima; construye 15 características."),
    ("2. Planificación y descomposición", "agent.py: trigger_time(), PLAN, step()", "Parcial", "Plan fijo de cinco pasos registrado en el log; no es un planificador dinámico."),
    ("3. Memoria", "memory.py: AircraftMemory, LongTermMemory", "Implementado", "Corto plazo: estado por aeronave. Largo plazo: historial simulado, climatología, tasas."),
    ("4. Razonamiento y decisión", "model.py: DelayModel; decision.py: decide()", "Implementado", "Modelo ML + función de utilidad. No se usa ningún LLM."),
    ("5. Llamada de herramientas y acción", "actuators.py: Actuators, OperationsSystemStub", "Simulado", "No hay integración con un sistema real de aerolínea ni con una API meteorológica real."),
    ("6. Comunicación", "actuators.py: send_alert()", "Implementado", "Mensaje estructurado (JSON) sin interfaz gráfica."),
    ("7. Aprendizaje y adaptación", "agent.py: receive_outcome(), retrain()", "Implementado", "Reentrenamiento por lotes al acumular 150 muestras (probado en S10); no hay reentrenamiento trimestral automático."),
], [4.3, 5.0, 2.2, 6.0], size=8)

doc.add_heading("9.2 Tecnologías", 2)
para("Qué razona el código y qué hace cada herramienta: la decisión (función de utilidad, validación, memoria, ciclo percepción → decisión → acción) la implementa el código propio; scikit-learn solo aporta el estimador de "
     "probabilidad y duración. No se usa ningún modelo de lenguaje.")
table(["Tecnología", "Versión", "Para qué se usó", "Por qué frente a la alternativa"], [
    ("Python", "3.12.3", "Lenguaje del agente, el simulador y las pruebas", "Ecosistema de ciencia de datos; alternativa Java, sin librerías ML equivalentes en el curso."),
    ("scikit-learn (HistGradientBoosting)", "1.8.0", "P(retraso) y duración esperada", "Admite valores faltantes y variables categóricas de forma nativa y entrena en CPU en segundos; la regresión logística no captura interacciones (hora × congestión) y una red neuronal exige más datos y ajuste."),
    ("NumPy / pandas", "2.4.4 / 3.0.2", "Simulación estocástica y manejo de tablas", "Estándar de facto; sin alternativas más simples para este volumen."),
    ("joblib", "1.5.3", "Guardar el modelo entrenado", "Formato recomendado por scikit-learn para modelos."),
    ("matplotlib / Graphviz", "3.10.8 / 2.43.0", "Figuras de resultados y diagramas", "Reproducibles desde script; las fuentes Mermaid (.mmd) y DOT (.dot) quedan en el repositorio."),
    ("python-docx", "ver requirements-doc.txt", "Generar este documento desde los resultados (scripts/construir_documento.py)", "Las cifras del texto salen de los CSV, lo que evita transcribir mal."),
    ("Función de utilidad propia", "—", "Decisión de alertar", "Hace explícito el compromiso entre falsas alarmas y retrasos anticipados; un umbral fijo no lo expresa."),
], [3.8, 2.4, 4.4, 6.9], size=8)

doc.add_heading("9.3 Entradas (inputs)", 2)
table(["Percepción", "Tipo", "Rango o formato", "Origen", "Ejemplo real (vuelo 1001-000192)"], [
    ("flight_id", "texto", "«semilla-número»", "Sistema de programación", "1001-000192"),
    ("tail (aeronave)", "texto", "identificador no vacío", "Sistema de programación", "T100100045"),
    ("leg (tramo)", "entero", "1 a 5", "Sistema de programación", "1"),
    ("airline", "categórico", "9 códigos (p. ej. AA, DL, UA, XP)", "Sistema de programación", "NK"),
    ("origin / dest", "categórico", "18 códigos IATA; origen ≠ destino", "Sistema de programación", "PHX / LAS"),
    ("sched_dep", "fecha-hora", "ISO 8601; posterior a «now»", "Sistema de programación", "2026-07-09T10:25:52"),
    ("prev_sched_arr", "fecha-hora o nulo", "ISO 8601", "Sistema de programación", "None"),
    ("now (reloj)", "fecha-hora", "ISO 8601; antelación = sched_dep − now ≥ 180 min", "Reloj del sistema", "2026-07-09T07:25:52"),
    ("forecast.vis", "decimal", "0 a 50 km", "Servicio meteorológico (simulado)", "9,7"),
    ("forecast.wind", "decimal", "0 a 200 kt", "Servicio meteorológico (simulado)", "7,6"),
    ("forecast.precip", "decimal", "0 a 500 mm", "Servicio meteorológico (simulado)", "0,0"),
    ("forecast.temp", "decimal", "−60 a 60 °C", "Servicio meteorológico (simulado)", "25,0"),
    ("Resultado real (retroalimentación)", "decimal + fecha-hora", "retraso en minutos; salida real", "Sistema operativo de la aerolínea (simulado)", "29,8 min"),
    ("Historial de entrenamiento", "tabla", "31.984 vuelos simulados, 20,1 % retrasados", "Simulador (sustituto del BTS)", "models/history_features.csv.gz"),
], [3.4, 2.4, 4.4, 3.8, 3.5], size=8)

doc.add_heading("9.4 Salidas y acciones", 2)
table(["Acción (mensaje)", "Cuándo se ejecuta", "Contenido", "Efecto en el entorno"], [
    ("ALERTA", "U(alertar) > 0", "flight_id, p_delay, d_hat_min, origin, tasa_historica_similares", "El equipo de operaciones conoce un riesgo que aún no se materializa."),
    ("SUGERENCIA_ REASIGNACION", "Junto con la alerta", "REASIGNAR_PUERTA o REASIGNAR_TRIPULACION_Y_PUERTA (si d ≥ 45 min o retraso en cadena estimado ≥ 30 min)", "El equipo puede reorganizar recursos antes del retraso."),
    ("REVISION_MANUAL", "Entrada inválida o antelación < 180 min", "flight_id, estado, problemas", "El vuelo queda pendiente de revisión humana; no se emite predicción."),
    ("REGISTRO_RESULTADO", "Al salir el vuelo", "flight_id, p_delay, alerta, retraso real", "Alimenta la memoria de la aeronave y el reentrenamiento."),
    ("Sin alerta (queda en el log)", "U(alertar) ≤ 0", "—", "Ninguno; se registra la decisión con su utilidad."),
], [3.6, 3.6, 5.6, 4.7], size=8)

doc.add_heading("9.5 Diagramas de secuencia UML", 2)
para("Participantes: actor externo (equipo de operaciones), entorno (simulador), interfaz (sistema operativo de la aerolínea), agente separado en planificación, percepción, memoria, decisión y acción, y el servicio externo (meteorológico). "
     "Los mensajes siguen el orden de las llamadas del código: trigger_time → perceive → extract_features → predict → decide → act → receive_outcome → retrain. Fuente editable en Mermaid: docs/diagramas/secuencia.mmd.")
figure(D / "diagramas" / "secuencia_tipico.png", 17.4, "Figura 2. Secuencia del escenario típico, con fragmentos loop, alt (anidado) y opt.")
figure(D / "diagramas" / "secuencia_excepcion.png", 17.4, "Figura 3. Secuencia de los escenarios de excepción: pronóstico ausente (imputación), entrada inválida y antelación insuficiente.")

doc.add_heading("9.6 Diferencias entre diseño e implementación", 2)
table(["Aspecto", "Diseño", "Implementación", "Por qué"], [
    ("Entorno y datos", "Historial BTS, API meteorológica y sistema operativo reales", "Simulador estocástico con verdad terreno oculta; pronóstico con error", "No hay acceso a datos ni sistemas reales; el simulador permite reproducibilidad y repeticiones."),
    ("Variables", "32 variables del BTS", "15 características (hora, día, mes, origen, destino, aerolínea, tramo, holgura, último retraso, retraso en cadena estimado, 4 variables de clima)", "Son las que el simulador puede generar y las que el diseño declara como entradas."),
    ("Regla de decisión", "Probabilidad contra un umbral definido", "Utilidad explícita; el umbral es P* = C_FA/(C_FA + β·d)", "Coherencia con un agente basado en utilidad."),
    ("Recálculo", "Continuo, a medida que se acerca el vuelo", "Una predicción a T − 180 min", "Alcance de la entrega; queda como trabajo futuro."),
    ("Estado de la aeronave", "Estado del vuelo anterior (a tiempo, retrasado, cancelado)", "Último retraso conocido en minutos; no se modelan cancelaciones", "Simplificación del simulador."),
    ("Aprendizaje", "Reentrenamiento periódico (trimestral)", "Reentrenamiento por lotes cada 150 resultados, con peso 5 para las muestras nuevas", "Permite probar el ciclo de aprendizaje en una sola ejecución (escenario S10)."),
    ("Actuadores", "Tres", "Cuatro (se agrega revisión manual)", "Cubre entradas inválidas y ventana insuficiente."),
    ("Medidas de desempeño", "Una medida compuesta", "Cuatro medidas (M1–M4)", "Ver tabla de cambios, sección 7."),
], [3.0, 4.0, 5.6, 4.9], size=8)

# ---- 9.7 pruebas --------------------------------------------------------------------------------------
doc.add_heading("9.7 Pruebas", 2)
n_runs, n_fl = res["total_corridas"], res["total_vuelos"]
para(f"**Método.** Cada medida se prueba en un entorno simulado con 13 escenarios (10 normales o límite y 3 de falla esperada); cada escenario se ejecuta 5 veces con semillas distintas "
     f"(semilla = 1000·n.º de escenario + repetición), es decir, {n_runs} corridas y {miles(n_fl)} vuelos. Los umbrales se fijaron en config/thresholds.json antes de ejecutar "
     f"({thr['fijado_utc'].replace('T', ' ').replace('Z', ' UTC')}) y no se modificaron. El modelo se entrenó con 31.984 vuelos simulados con semilla distinta de la de las pruebas; "
     "el parámetro β = 0,05 se definió en el diseño y no se ajustó después de ver los resultados. "
     "**Reproducibilidad:** los resultados de esta tabla se obtuvieron con las versiones de 9.2. Una reejecución independiente con Python 3.13 y scikit-learn 1.9.1 dio M1 = 74,0 %, M2 = 100 %, "
     "M4 = 100 % y M3 = 46,2 % (en lugar de 46,0 %): las conclusiones no cambian, pero las décimas varían con la versión de las librerías.")
doc.add_heading("Tabla de resultados", 3)
cumple = lambda b: "SÍ" if b else "NO"
rows = []
for _, r_ in tabla.iterrows():
    rows.append((r_["medida"], r_["formula"].replace(">=", "≥").replace("VP+VN", "VP + VN").replace("VP/(VP+FN)", "VP / (VP + FN)"), r_["umbral"].replace(">=", "≥").replace("%", " %").replace("  ", " "),
                 r_["casos"].replace(" x ", " × "), r_["resultado"].replace(".", ","), cumple(r_["cumple"]), r_["evidencia"]))
table(["Medida de desempeño", "Fórmula", "Umbral", "Casos ejecutados", "Resultado obtenido", "¿Cumple?", "Evidencia (log, archivo)"], rows,
      [2.9, 2.9, 1.5, 2.6, 2.4, 1.7, 3.6], size=7.5)

doc.add_heading("Escenarios de prueba (13): normales, límite y fallas esperadas", 3)
srows = []
for _, e in esc.iterrows():
    info = e["tipo"] == "falla"
    srows.append((e["escenario"], e["nombre"].replace("FALLA: ", ""), {"limite": "límite"}.get(e["tipo"], e["tipo"]),
                  f"{pc(e['exactitud_media'])} ± {num(e['exactitud_std'] * 100)}" + (" (inf.)" if info else ""),
                  f"{pc(e['recall_media'])} ± {num(e['recall_std'] * 100)}" + (" (inf.)" if info else ""),
                  pc(e["baseline_acc"])))
table(["ID", "Escenario", "Tipo", "Exactitud media ± DE", "Recall media ± DE", "Línea base (nunca alertar)"], srows,
      [1.1, 6.0, 1.7, 3.2, 3.2, 2.8], size=8)
para("(inf.) = informativa: en los escenarios de falla la medida que se evalúa es M4, no M1/M3.", italic=True, size=8.5)
figure(R / "fig_resultados_escenarios.png", 15.0, "Figura 4. Exactitud y recall por escenario (media ± desviación estándar sobre 5 repeticiones). Los marcadores rojos son la exactitud de la línea base «nunca alertar».")
figure(R / "fig_compromiso_beta.png", 12.0, "Figura 5. Compromiso exactitud–recall al variar β (datos de validación). La línea vertical marca el valor usado (β = 0,05).")

# ---- análisis (<= 1 página) --------------------------------------------------------------------------
m1, m3 = res["M1_media"], res["M3_media"]
below = esc[(esc.tipo != "falla") & (esc.exactitud_media < 0.70)]
worse = int(((esc.tipo != "falla") & (esc.exactitud_media < esc.baseline_acc)).sum())
doc.add_heading("Análisis de resultados", 3)
para(f"**M1 (exactitud): cumple, con una advertencia.** La media global es {pc(m1)} (DE {num(res['M1_std'] * 100)} puntos; umbral 70 %). Tres escenarios quedan por debajo de 70 % por separado: "
     + ", ".join(f"{e['escenario']} ({pc(e['exactitud_media'])})" for _, e in below.iterrows()) +
     f". La advertencia: en {worse} de los 10 escenarios la exactitud del agente es menor que la de la línea base «nunca alertar» (p. ej. S01: {pc(esc.iloc[0]['exactitud_media'])} frente a {pc(esc.iloc[0]['baseline_acc'])}), "
     "porque con ≈ 20 % de retrasos esa línea base ya supera el umbral. M1 se cumple, pero por sí sola no demuestra que el agente aporta valor; por eso se añadió M3.")
para("**M2 (antelación): cumple (100 %).** Ninguna de las predicciones emitidas tuvo menos de 180 min. Es cierto por construcción del disparador T − 180 y por la validación de ventana: "
     "en S12 los registros tardíos se derivaron a revisión manual en lugar de predecirse. Con datos reales, la antelación dependerá también de cuándo llegue el registro desde el sistema de programación.")
para(f"**M3 (recall): NO cumple, {pc(m3)} frente a un umbral de 60 %.** La dispersión entre escenarios es grande (de {pc(esc[esc.tipo != 'falla'].recall_media.min())} en la madrugada a {pc(esc[esc.tipo != 'falla'].recall_media.max())} en tormenta). "
     "Causas probables: (1) con β = 0,05 y C_FA = 1 el agente alerta solo si P supera ≈ 0,25–0,40 y el modelo rara vez llega a esa probabilidad en retrasos «ordinarios»; "
     "(2) observabilidad parcial: el retraso del tramo previo solo se conoce si ese vuelo ya salió a T − 180 min y los incidentes ocultos son impredecibles; (3) el pronóstico tiene error. "
     "Evidencia en validación (datos distintos de los de prueba): "
     + "; ".join(f"con β = {num(s_['BETA'], 2)} el recall es {pc(s_['recall'])} y la exactitud {pc(s_['exactitud'])}" for _, s_ in sens[sens.BETA.round(2).isin([0.08, 0.10])].iterrows()) +
     ". Con este modelo, M1 ≥ 70 % y M3 ≥ 60 % no pueden cumplirse a la vez: hay un compromiso entre falsas alarmas y retrasos detectados (Figura 5). "
     "**Qué cambiaríamos:** (a) dar más información al modelo (recalcular a T − 90 y T − 60 min con el estado del vuelo previo ya conocido y usar datos reales con más variables); "
     "(b) acordar con operaciones cuánto cuesta una falsa alarma frente a un retraso no anticipado y fijar β y C_FA en consecuencia; "
     "(c) reformular M1 y M3 como una sola medida (F1 o costo esperado) con un umbral nuevo justificado antes de volver a probar.")
para("**M4 (manejo de fallas): cumple (100 % en S11, S12 y S13, 0 excepciones).** Los registros corruptos y los tardíos se marcaron para revisión manual sin marcar ningún registro válido por error, "
     "y los vuelos con pronóstico ausente recibieron una predicción con clima imputado y bandera. "
     "**Aprendizaje (S10):** el agente se reentrenó 3 o 4 veces por corrida; la diferencia entre el primer y el último cuarto del flujo es del orden de la desviación entre repeticiones, "
     "por lo que no se puede afirmar que el reentrenamiento mejore el desempeño, solo que el ciclo de retroalimentación funciona de extremo a extremo.")

doc.add_heading("9.8 Limitaciones y trabajo futuro", 2)
for b in ("Los resultados son contra un simulador escrito por el equipo: reflejan la lógica del agente, no su desempeño en el tráfico aéreo real. Con datos del BTS las cuatro medidas deben repetirse.",
          "El simulador se calibró para una tasa de retrasos cercana al 20 % (la tasa histórica típica del BTS); si es más o menos predecible que la realidad, M1 y M3 cambiarían.",
          "Una sola predicción a T − 180 min; falta el recálculo a medida que se acerca el vuelo y modelar cancelaciones.",
          "Sin integración real con un sistema operativo de aerolínea ni con una API meteorológica; sin interfaz gráfica (la demo es por consola).",
          "Trabajo futuro: reevaluación periódica, función de utilidad calibrada con costos reales, calibración de probabilidades, análisis de equidad por aeropuerto y aerolínea, y despliegue en una plataforma pública."):
    bullet(b)

doc.add_heading("9.9 Enlaces", 2)
table(["Recurso", "Enlace"], [
    ("Repositorio (GitHub/GitLab)", "[[PEGAR enlace del repositorio con README y requirements.txt]]"),
    ("Video (≤ 20 min) o plataforma de despliegue", "[[PEGAR enlace del video con índice de minutos, o de la plataforma desplegada]]"),
    ("Logs de pruebas", "results/logs/*.jsonl.gz (copias sin comprimir en results/logs_muestra/), results/vuelos_detalle.csv.gz, results/detalle_por_corrida.csv, results/tabla_resultados.csv"),
], [5.5, 12.0], size=9)

# =====================================================================================================
doc.add_heading("10. Conclusiones", 1)
para(f"FlightPredict-Agent se implementó como un agente basado en utilidad con modelo interno y aprendizaje, con un ciclo percepción → decisión → acción identificable en el código, memoria de corto y largo plazo, "
     f"un registro de ejecución con marca de tiempo y un entorno simulado que respeta las propiedades de ambiente del diseño. Las pruebas, con umbrales fijados antes de ejecutar, dieron tres medidas cumplidas "
     f"(exactitud de {pc(m1)}, antelación de 100 % y manejo de fallas de 100 %) y una incumplida (recall de {pc(m3)} frente a 60 %).")
para("El hallazgo principal es que la medida de exactitud original era insuficiente: con ≈ 20 % de vuelos retrasados, la línea base «nunca alertar» ya la supera, y subir la sensibilidad del agente cuesta exactitud. "
     "El diseño debe, entonces, definir con el equipo de operaciones el costo relativo de una falsa alarma y de un retraso no anticipado, y medir el agente con una métrica que los combine. "
     "Desde la perspectiva ética se mantiene lo planteado en el diseño: el agente es un apoyo a la decisión, supervisado por humanos, y debe revisarse el sesgo por aeropuerto y aerolínea antes de usarse con datos reales.")

doc.add_heading("Referencias", 1)
for b in ("Russell, S. J., & Norvig, P. Artificial Intelligence: A Modern Approach (Global Edition).",
          "IBM. «Components of AI agents». https://www.ibm.com/mx-es/think/topics/components-of-ai-agents",
          "Bureau of Transportation Statistics (BTS), U.S. Department of Transportation. Datos históricos de vuelos domésticos. https://www.transtats.bts.gov",
          "«Flight delay prediction: Evaluating machine learning algorithms for enhanced accuracy». PLOS ONE. https://journals.plos.org/plosone/article?id=10.1371/journal.pone.0335141",
          "Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
          "Hendrycks, D., Burns, C., Chen, A., & Ball, S. (2021). CUAD: An Expert-Annotated NLP Dataset for Legal Contract Review. arXiv:2103.06268."):
    bullet(b)

# ---- Anexo A: extracto real del log ---------------------------------------------------------------------
doc.add_heading("Anexo A. Extracto del registro de ejecución (S01, repetición 1, vuelo 1001-000192)", 1)
para("Cada línea del log (JSONL comprimido) guarda la marca de tiempo real (ts_wall), la marca de tiempo simulada (ts_sim), la fase y los datos. Se muestran la percepción, la decisión, la acción y la retroalimentación del mismo vuelo que se usa como ejemplo en 9.3 (campos largos abreviados con «…»).", size=9.5)
lines = []
for l in gzip.open(R / "logs" / "S01_rep1.jsonl.gz", "rt"):
    o = json.loads(l)
    if o["flight_id"] == "1001-000192":
        d_ = o["data"]
        if "plan" in d_:
            d_ = {**d_, "plan": "[…5 pasos…]"}
        d_ = {k: (round(v, 3) if isinstance(v, float) else v) for k, v in d_.items()}
        lines.append(f'{{"ts_wall": "{o["ts_wall"][:23]}", "ts_sim": "{o["ts_sim"]}", "phase": "{o["phase"]}", "data": {json.dumps(d_, ensure_ascii=False)}}}')
for ln in lines:
    p = doc.add_paragraph()
    r_ = p.add_run(ln); r_.font.name = "Consolas"; r_.font.size = Pt(7)
    p.paragraph_format.space_after = Pt(3)

doc.save(OUT)
print("OK ->", OUT)
